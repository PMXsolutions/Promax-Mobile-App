#!/usr/bin/env python3
"""Generate Promax Care Mobile App Testing / Bug Report Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCS_DIR = Path(__file__).resolve().parent
SCREENSHOTS_DIR = DOCS_DIR / "screenshots"

FIGURES = {
    "fig-01": {
        "file": "fig-01-shift-detail-not-started.jpg",
        "id": "Figure 1",
        "title": "Staff shift detail — Not Started, transport control visible",
        "caption": (
            "Staff mobile shift detail for client Ajeh Wumi Tao (staff Serah Adeniyi). "
            "Date shows 3 October 2026, start 4:15 PM / end 5:18 PM. Status is “Not Started” "
            "with message “Your shift hasn't started yet.” Transport (steering wheel) button "
            "is still visible. CTA is “Request to Cancel Shift”; clock-in is not offered."
        ),
        "width_in": 2.35,
    },
    "fig-02": {
        "file": "fig-02-admin-schedule-board.jpg",
        "id": "Figure 2",
        "title": "Admin Schedule Board — same shift In Progress (Sydney)",
        "caption": (
            "ProMax Care web Schedule Board at app.promaxcare.com.au (header 4:17:31 pm – Sydney, AU). "
            "Saturday 3 October 2026, 4:15 PM – 5:18 PM for staff Serah Adeniyi / client Ajeh Wumi Tao "
            "is shown as “In Progress” with activities Transport, Install phone, Personal Support."
        ),
        "width_in": 3.6,
    },
    "fig-03": {
        "file": "fig-03-staff-roster-4-oct.jpg",
        "id": "Figure 3",
        "title": "Staff roster dashboard — shift listed on 4 October",
        "caption": (
            "Staff mobile Shift Roster for Serah. Calendar selects Sunday 4 October 2026. "
            "The same client/activities card shows 1:15 AM – 3:18 AM with status ACTIVE, "
            "instead of Saturday 3 October 4:15 PM – 5:18 PM Australian time."
        ),
        "width_in": 2.35,
    },
    "fig-04": {
        "file": "fig-04-add-document-no-images.jpg",
        "id": "Figure 4",
        "title": "Add a document — file types limited to PDF and Word",
        "caption": (
            "Staff “Add a document” screen. The Upload File control reads "
            "“Tap to select file (.pdf, .doc, .docx)”. Image formats (JPEG, PNG, HEIC) "
            "are not listed or offered."
        ),
        "width_in": 2.35,
    },
    "fig-05": {
        "file": "fig-05-edit-report-two-clients.jpg",
        "id": "Figure 5",
        "title": "Edit Shift Report — two clients on one report, no selector",
        "caption": (
            "Edit Shift Report for staff Serah. Client’s Name is shown as a combined string "
            "“Ajeh Wumi Tao, Olatunji John”. There is no control to choose which client the "
            "report is for. One form (urgent matters, medications, etc.) applies to both clients."
        ),
        "width_in": 2.35,
    },
    "fig-06": {
        "file": "fig-06-completed-shift-multi-client-trip.jpg",
        "id": "Figure 6",
        "title": "Completed multi-client shift — trip control still available",
        "caption": (
            "Shift detail for clients “David David David, Fateru Israel Oluwapelumi” "
            "(staff Serah Adeniyi). Status is “Shift Completed” with message "
            "“Great job! Your shift has been completed successfully.” Activities include "
            "Transport. The steering-wheel trip button remains visible. There is no client "
            "selector for which participant a trip would cover."
        ),
        "width_in": 2.35,
    },
}


def set_cell_shading(cell, color_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x03, 0x06, 0x37)
    return heading


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "030637")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)

    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = str(cell_text)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph()
    return table


def severity_para(doc, label, color_hex):
    p = doc.add_paragraph()
    run = p.add_run(f"Severity: {label}")
    run.bold = True
    run.font.color.rgb = RGBColor(
        int(color_hex[0:2], 16),
        int(color_hex[2:4], 16),
        int(color_hex[4:6], 16),
    )
    return p


def add_figure(doc, figure_key: str):
    """Embed a screenshot and caption. Returns True if the image was inserted."""
    fig = FIGURES[figure_key]
    path = SCREENSHOTS_DIR / fig["file"]

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = cap.add_run(f"{fig['id']} — {fig['title']}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x03, 0x06, 0x37)

    if not path.exists():
        missing = doc.add_paragraph()
        missing.add_run(
            f"[Screenshot missing: {fig['file']}. Place the file in docs/screenshots/ and regenerate.]"
        ).italic = True
        return False

    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    run.add_picture(str(path), width=Inches(fig["width_in"]))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = caption.add_run(fig["caption"])
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    caption.paragraph_format.space_after = Pt(12)
    return True


def build_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Promax Care Mobile App")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x03, 0x06, 0x37)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Testing Report — Staff Shift Issues")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Report Date: {date.today().strftime('%d %B %Y')}  |  App Version: 1.1.0  |  Status: Open  |  Format: Microsoft Word (.docx)"
    )
    run.font.size = Pt(10)
    run.italic = True

    # 1. Summary
    add_heading(doc, "1. Summary", 1)
    doc.add_paragraph(
        "This short testing report documents four defects observed during staff shift operations "
        "on the Promax Care mobile application. Issues were reproduced against a live shift for "
        "staff Serah Adeniyi and client Ajeh Wumi Tao (3 October 2026, 4:15 PM – 5:18 PM Australian time). "
        "Evidence is attached as embedded screenshots (Figures 1–6) captured from the staff mobile app "
        "and the admin Schedule Board."
    )

    add_table(
        doc,
        ["ID", "Issue", "Severity", "Status", "Screenshot refs"],
        [
            [
                "BUG-01",
                "Trip can start before/after shift window",
                "High",
                "Open",
                "Figures 1, 6",
            ],
            [
                "BUG-02",
                "Shift date/time shown on wrong day; clock-in blocked",
                "Critical",
                "Open",
                "Figures 1, 2, 3",
            ],
            [
                "BUG-03",
                "Document upload rejects images",
                "High",
                "Open",
                "Figure 4",
            ],
            [
                "BUG-04",
                "No client selection for multi-client report/trip",
                "High",
                "Open",
                "Figures 5, 6",
            ],
        ],
    )

    # 2. Test Environment
    add_heading(doc, "2. Test Environment", 1)
    add_table(
        doc,
        ["Item", "Detail"],
        [
            ["Application", "Promax Care Mobile (Staff)"],
            ["Admin Reference", "ProMax Care web Schedule Board (app.promaxcare.com.au)"],
            ["Business Timezone", "Australia/Sydney"],
            ["Test Staff", "Serah Adeniyi"],
            ["Test Client", "Ajeh Wumi Tao"],
            ["Test Shift", "Sat 3 Oct 2026, 4:15 PM – 5:18 PM (Australian time)"],
            ["Activities", "Transport, Install phone, Personal Support"],
            ["Admin Board Status (at ~4:17 PM Sydney)", "In Progress — see Figure 2"],
            ["Staff Detail Status (device ~7:20)", "Not Started — see Figure 1"],
            ["Staff Roster Display", "Sunday 4 Oct, 1:15 AM – 3:18 AM, ACTIVE — see Figure 3"],
        ],
    )

    add_heading(doc, "2.1 Screenshot Catalogue", 2)
    add_table(
        doc,
        ["Ref", "Source", "What it shows"],
        [
            [
                "Figure 1",
                "Staff mobile — shift detail",
                "Date 3 Oct; times 4:15–5:18 PM; Not Started; transport button visible; no clock-in",
            ],
            [
                "Figure 2",
                "Admin web — Schedule Board",
                "Same shift Sat 3 Oct 4:15–5:18 PM Sydney shown In Progress",
            ],
            [
                "Figure 3",
                "Staff mobile — Shift Roster",
                "Same shift listed on Sunday 4 Oct at 1:15–3:18 AM, ACTIVE",
            ],
            [
                "Figure 4",
                "Staff mobile — Add a document",
                "Upload File accepts only .pdf, .doc, .docx — no images",
            ],
            [
                "Figure 5",
                "Staff mobile — Edit Shift Report",
                "Client’s Name: Ajeh Wumi Tao, Olatunji John — one report, no picker",
            ],
            [
                "Figure 6",
                "Staff mobile — shift detail (completed)",
                "Two clients combined; Shift Completed; transport button still shown",
            ],
        ],
    )
    doc.add_paragraph(
        "Screenshots are embedded in the relevant defect sections below and repeated in Appendix A."
    )

    # BUG-01
    add_heading(doc, "3. BUG-01 — Travel Trip Available Outside Shift Window", 1)
    severity_para(doc, "High", "C0392B")
    p = doc.add_paragraph()
    run = p.add_run("Screenshot reference: ")
    run.bold = True
    p.add_run("Figure 1 (before start); Figure 6 (after shift completed)")

    add_heading(doc, "3.1 Description", 2)
    doc.add_paragraph(
        "Staff can start or process a travelling trip before the shift has begun and after the "
        "shift has ended. Trip/travel actions should only be available while the shift is active "
        "(clocked in / in progress)."
    )

    add_heading(doc, "3.2 Evidence", 2)
    add_bullet(
        doc,
        " Shift detail shows status “Not Started” and “Your shift hasn't started yet. "
        "You can request to cancel if needed.” The transport (steering wheel) control remains "
        "visible and usable (Figure 1).",
        "Before start: ",
    )
    add_bullet(
        doc,
        " A later completed shift for two clients still shows “Shift Completed” while the "
        "transport button remains on screen (Figure 6).",
        "After end: ",
    )
    add_bullet(
        doc,
        " Transport button is shown whenever the shift activities include “Transport”, with no "
        "check that the shift is in progress. Status-based guards are commented out in code "
        "(shift detail screen).",
        "Code note: ",
    )
    add_figure(doc, "fig-01")
    add_figure(doc, "fig-06")

    add_heading(doc, "3.3 Steps to Reproduce", 2)
    for step in [
        "Open a shift that includes the Transport activity (see Figure 1 before start, Figure 6 after complete).",
        "Confirm shift status is Not Started / Upcoming (before start), or Shift Completed / Absent / Present (after end).",
        "Tap the transport / trip control (steering-wheel button in Figures 1 and 6).",
        "Observe that trip tracking can still be started or processed.",
    ]:
        doc.add_paragraph(step, style="List Number")

    add_heading(doc, "3.4 Expected vs Actual", 2)
    add_table(
        doc,
        ["Expected", "Actual"],
        [
            [
                "Trip start/process only when shift is In Progress (staff clocked in).",
                "Trip control available before start (Figure 1) and after the shift is completed (Figure 6).",
            ]
        ],
    )

    # BUG-02
    add_heading(doc, "4. BUG-02 — Incorrect Shift Date/Time and Blocked Clock-In", 1)
    severity_para(doc, "Critical", "8B0000")
    p = doc.add_paragraph()
    run = p.add_run("Screenshot references: ")
    run.bold = True
    p.add_run("Figure 2 (admin), Figure 3 (staff roster), Figure 1 (staff detail)")

    add_heading(doc, "4.1 Description", 2)
    doc.add_paragraph(
        "A shift created for 3 October 2026, 4:15 PM – 5:18 PM Australian time appears on the "
        "staff dashboard as 4 October with times 1:15 AM – 3:18 AM. Opening shift detail shows "
        "Date as 3 October, but status remains “Not Started” and clock-in is unavailable, even "
        "when the admin Schedule Board (Sydney time) shows the same shift as “In Progress”."
    )
    doc.add_paragraph(
        "This indicates inconsistent timezone handling between roster day grouping, detail date "
        "display, clock-in window logic, and/or device local time versus Australia/Sydney."
    )

    add_heading(doc, "4.2 Evidence", 2)
    add_bullet(
        doc,
        " Admin Schedule Board at 4:17:31 pm – Sydney, AU shows shift Sat Oct 3, "
        "4:15 PM – 5:18 PM as “In Progress” (Figure 2).",
        "Admin: ",
    )
    add_bullet(
        doc,
        " Staff Shift Roster selects Sunday Oct 4 and shows the same client/activities at "
        "1:15 AM – 3:18 AM with status ACTIVE (Figure 3).",
        "Staff roster: ",
    )
    add_bullet(
        doc,
        " Shift detail Date field is “3 October, 2026”; start/end 4:15 PM / 5:18 PM; "
        "status “Not Started”; CTA “Request to Cancel Shift” with no clock-in (Figure 1).",
        "Staff detail: ",
    )
    add_bullet(
        doc,
        " Detail “Date” uses dateCreated (not dateFrom). Roster calendar day chips use "
        "device-local dates while shift grouping/times use Australia/Sydney helpers — "
        "mixed sources can place a shift on the wrong day and block clock-in.",
        "Code note: ",
    )

    add_heading(doc, "4.2.1 Admin (expected source of truth)", 3)
    add_figure(doc, "fig-02")

    add_heading(doc, "4.2.2 Staff roster (wrong calendar day / converted times)", 3)
    add_figure(doc, "fig-03")

    add_heading(doc, "4.2.3 Staff detail (date 3 Oct, clock-in blocked)", 3)
    add_figure(doc, "fig-01")

    add_heading(doc, "4.3 Steps to Reproduce", 2)
    for step in [
        "In admin, create/confirm a shift for 3 Oct 2026, 4:15 PM – 5:18 PM Australian time (Figure 2).",
        "On staff mobile, open Shift Roster and locate the shift (Figure 3).",
        "Observe which calendar day and start/end times are shown.",
        "Open shift detail and compare Date, Start Time, End Time, and status (Figure 1).",
        "Attempt to clock in while admin board shows the shift In Progress in Sydney time.",
    ]:
        doc.add_paragraph(step, style="List Number")

    add_heading(doc, "4.4 Expected vs Actual", 2)
    add_table(
        doc,
        ["Expected", "Actual", "Evidence"],
        [
            [
                "Roster and detail show 3 Oct 2026, 4:15 PM – 5:18 PM (Sydney).",
                "Roster shows 4 Oct, 1:15 AM – 3:18 AM; detail Date shows 3 Oct.",
                "Figures 3 and 1 vs Figure 2",
            ],
            [
                "Staff can clock in once Sydney window is open / shift has started.",
                "Status stays Not Started; clock-in unavailable; only cancel offered.",
                "Figure 1",
            ],
            [
                "Staff status matches admin Schedule Board for the same shift.",
                "Admin: In Progress; Staff detail: Not Started.",
                "Figures 2 and 1",
            ],
        ],
    )

    # BUG-03
    add_heading(doc, "5. BUG-03 — Document Upload Does Not Accept Images", 1)
    severity_para(doc, "High", "C0392B")
    p = doc.add_paragraph()
    run = p.add_run("Screenshot reference: ")
    run.bold = True
    p.add_run("Figure 4")

    add_heading(doc, "5.1 Description", 2)
    doc.add_paragraph(
        "Staff document upload should accept image files (e.g. photos of certificates and IDs). "
        "In practice, image selection/upload is not accepted as expected for compliance documents."
    )

    add_heading(doc, "5.2 Evidence", 2)
    add_bullet(
        doc,
        " The “Add a document” screen labels Upload File as “Tap to select file (.pdf, .doc, .docx)”. "
        "No image types are listed (Figure 4).",
        "Observed: ",
    )
    add_bullet(
        doc,
        " Upload uses DocumentPicker with a limited MIME list: image/jpeg, application/pdf, "
        "and Word types only. The on-screen hint omits images entirely. image/png and camera/gallery "
        "paths are not offered (unlike profile photo).",
        "Code / UX note: ",
    )
    add_figure(doc, "fig-04")

    add_heading(doc, "5.3 Steps to Reproduce", 2)
    for step in [
        "Open Documents tab → Add a document (Figure 4).",
        "Confirm the upload hint lists only .pdf, .doc, .docx.",
        "Tap “Tap to select file” and attempt to choose a JPEG, PNG, HEIC, or camera photo.",
        "Observe that images are not offered or cannot be submitted as a compliance document.",
    ]:
        doc.add_paragraph(step, style="List Number")

    add_heading(doc, "5.4 Expected vs Actual", 2)
    add_table(
        doc,
        ["Expected", "Actual", "Evidence"],
        [
            [
                "Staff can upload common image formats (JPEG, PNG; ideally HEIC via conversion) "
                "and PDF/Word for compliance documents.",
                "Upload File hint allows only .pdf, .doc, .docx. Images are not accepted.",
                "Figure 4",
            ]
        ],
    )

    # BUG-04
    add_heading(doc, "6. BUG-04 — No Client Selection for Multi-Client Report / Trip", 1)
    severity_para(doc, "High", "C0392B")
    p = doc.add_paragraph()
    run = p.add_run("Screenshot references: ")
    run.bold = True
    p.add_run("Figure 5 (report); Figure 6 (trip on completed multi-client shift)")

    add_heading(doc, "6.1 Description", 2)
    doc.add_paragraph(
        "When multiple clients are assigned to a staff member on a shift, a single shift report "
        "and a single travelled trip are tied to both clients together. Staff cannot choose "
        "which client the report applies to, or which client a trip covers."
    )

    add_heading(doc, "6.2 Evidence", 2)
    add_bullet(
        doc,
        " Edit Shift Report shows Client’s Name as “Ajeh Wumi Tao, Olatunji John” with no "
        "dropdown or picker. One report form is shared by both clients (Figure 5).",
        "Report: ",
    )
    add_bullet(
        doc,
        " Shift detail header concatenates two clients (“David David David, Fateru Israel "
        "Oluwapelumi”). Transport remains a single shift-level button even after completion "
        "(Figure 6). There is no client selector for the trip.",
        "Trip: ",
    )
    add_bullet(
        doc,
        " Shift model exposes a combined clients string and a single profile object; "
        "report submit payload sends shiftRosterId + form fields with no clientId selector.",
        "Data model: ",
    )

    add_heading(doc, "6.2.1 One report tied to two clients", 3)
    add_figure(doc, "fig-05")

    add_heading(doc, "6.2.2 One trip control for two clients (after shift completed)", 3)
    add_figure(doc, "fig-06")

    add_heading(doc, "6.3 Steps to Reproduce", 2)
    for step in [
        "Assign a shift with more than one client for the same staff member.",
        "Open Edit/Create Shift Report and confirm Client’s Name lists both clients with no selector (Figure 5).",
        "Submit or save the report and observe it is stored against the shift, not a chosen client.",
        "Open shift detail for a multi-client shift that includes Transport (Figure 6).",
        "Start/save a trip via the steering-wheel control. Observe the trip is not attributed to a specific client.",
    ]:
        doc.add_paragraph(step, style="List Number")

    add_heading(doc, "6.4 Expected vs Actual", 2)
    add_table(
        doc,
        ["Expected", "Actual", "Evidence"],
        [
            [
                "Staff selects which client a shift report is for when multiple clients apply.",
                "One report is tied to all clients on the shift (combined name string).",
                "Figure 5",
            ],
            [
                "Staff selects which client a travelled trip covers.",
                "Trip is recorded against the shift only; transport button is shift-level.",
                "Figure 6",
            ],
        ],
    )

    # Recommendations
    add_heading(doc, "7. Recommended Fixes (Testing Focus)", 1)
    add_table(
        doc,
        ["Bug", "Recommended Fix Direction", "Suggested Retest"],
        [
            [
                "BUG-01",
                "Gate TransportButton / trip start to Clock-In or Shift In Progress only; "
                "block after Present/Absent/Cancelled/Upcoming.",
                "Before start, during active, after end — only mid-shift allows trip. Compare to Figure 1.",
            ],
            [
                "BUG-02",
                "Standardise all roster day keys, calendar selection, detail Date, and "
                "clock-in windows on Australia/Sydney; show dateFrom (not dateCreated) as "
                "shift date; align staff status with admin Sydney clock.",
                "Create Sydney-afternoon shift; verify day, times, In Progress, and clock-in "
                "on a device set to a non-AU timezone. Compare Figures 1–3.",
            ],
            [
                "BUG-03",
                "Accept image/jpeg, image/png, and preferably ImagePicker/camera with "
                "conversion; keep PDF/Word support.",
                "Upload JPEG, PNG, and camera photo for a required document type. Compare to Figure 4.",
            ],
            [
                "BUG-04",
                "Add mandatory client selector on report create and trip start when "
                "clients.length > 1; persist clientId on report and trip payloads.",
                "Multi-client shift: submit two reports and two trips for different clients. Compare to Figures 5 and 6.",
            ],
        ],
    )

    # Appendix
    add_heading(doc, "8. Appendix A — Embedded Screenshots", 1)
    doc.add_paragraph(
        "The following figures are the original test captures. They are also placed inline "
        "under BUG-01 through BUG-04."
    )
    add_figure(doc, "fig-01")
    add_figure(doc, "fig-02")
    add_figure(doc, "fig-03")
    add_figure(doc, "fig-04")
    add_figure(doc, "fig-05")
    add_figure(doc, "fig-06")

    # Sign-off
    add_heading(doc, "9. Sign-Off", 1)
    add_table(
        doc,
        ["Role", "Name", "Date", "Signature"],
        [
            ["Tester", "", "", ""],
            ["Developer", "", "", ""],
            ["Product / Ops", "", "", ""],
        ],
    )

    footer = doc.add_paragraph()
    run = footer.add_run(
        "Document format: Microsoft Word (.docx). Embedded files: "
        "fig-01-shift-detail-not-started.jpg, fig-02-admin-schedule-board.jpg, "
        "fig-03-staff-roster-4-oct.jpg, fig-04-add-document-no-images.jpg, "
        "fig-05-edit-report-two-clients.jpg, fig-06-completed-shift-multi-client-trip.jpg "
        "(docs/screenshots/)."
    )
    run.italic = True
    run.font.size = Pt(9)

    return doc


if __name__ == "__main__":
    output_path = DOCS_DIR / "Promax_Care_Testing_Report.docx"
    build_document().save(output_path)
    print(f"Document saved to: {output_path}")
